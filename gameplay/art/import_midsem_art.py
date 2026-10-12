"""Use the Art Team's *original full-resolution* images from the MidSem deck.

P1 idle is extracted byte-for-byte from slide 13 (no crop / resize).
P2 idle is a pixel-for-pixel copy with only blue costume RGB recolored red.
Both animation sets retain the full original 736x736 GIF frame geometry.
All display scaling happens in Godot's Sprite2D, not these source files.
"""
from pathlib import Path
from io import BytesIO
from PIL import Image
from colorsys import rgb_to_hsv, hsv_to_rgb
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
DECK = Path('/Users/suwen/Downloads/Property Damage MidSem Presentation (1).pptx')
ASSET = ROOT / 'gameplay/art'
SCENE_ASSET = ROOT / 'gameplay/msd/art'
P = Presentation(str(DECK))


def slide_blob(slide: int, image_index: int) -> bytes:
    imgs = [shape.image.blob for shape in P.slides[slide].shapes if shape.shape_type == 13]
    return imgs[image_index]


def to_red(source: Image.Image) -> Image.Image:
    """Recolor blues only; preserve every original alpha and non-blue pixel."""
    src = source.convert('RGBA')
    dst = src.copy()
    original = src.load()
    colored = dst.load()
    for y in range(src.height):
        for x in range(src.width):
            r, g, b, a = original[x, y]
            if not a:
                continue
            h, saturation, value = rgb_to_hsv(r / 255, g / 255, b / 255)
            # Select saturated blue shades of the original robe/hat only.
            if 0.51 <= h <= 0.76 and saturation > 0.21 and b >= max(r, g) * 0.99:
                rr, gg, bb = hsv_to_rgb(0.0, min(0.95, saturation * 0.98), value)
                colored[x, y] = (round(rr * 255), round(gg * 255), round(bb * 255), a)
    return dst


static_png = slide_blob(12, 0)
static = Image.open(BytesIO(static_png)).convert('RGBA')
original_gif = slide_blob(12, 2)
run_gif = Image.open(BytesIO(original_gif))
assert static.size == (736, 736) and run_gif.size == (736, 736) and run_gif.n_frames == 2

# P1 is the exact original PNG and GIF bytestream, including the transparent
# margins. P2 has identical coordinates, silhouette and transparency.
(ASSET / 'mage_1.png').write_bytes(static_png)
(ASSET / 'mage_1_run.gif').write_bytes(original_gif)
to_red(static).save(ASSET / 'mage_2.png', optimize=True)
red_frames = []
for index in range(2):
    run_gif.seek(index)
    blue_frame = run_gif.convert('RGBA')
    blue_frame.save(ASSET / f'mage_1_run_{index}.png', optimize=True)
    red_frame = to_red(blue_frame)
    red_frame.save(ASSET / f'mage_2_run_{index}.png', optimize=True)
    red_frames.append(red_frame)
# Preview GIF; Godot uses full-resolution PNG frames instead of decoding GIFs.
red_frames[0].save(
    ASSET / 'mage_2_run.gif', save_all=True, append_images=red_frames[1:],
    duration=[80, 80], loop=0, disposal=2, optimize=False,
)

# Keep the previously imported Art Team furniture sprites reproducible.
collage = Image.open(BytesIO(slide_blob(11, 0))).convert('RGBA')
assert collage.size == (529, 384)
regions = {
    'team_shelf': (0, 154, 178, 384),
    'team_ac': (168, 0, 400, 70),
    'team_monitor': (234, 183, 359, 281),
    'team_fridge': (358, 202, 464, 384),
    'team_window': (0, 0, 134, 103),
    'team_clock': (446, 0, 529, 83),
}
for name, rect in regions.items():
    collage.crop(rect).save(SCENE_ASSET / f'{name}.png', optimize=True)

# Byte and pixel identity checks catch accidental rescaling or editing.
assert (ASSET / 'mage_1.png').read_bytes() == static_png
red_idle = Image.open(ASSET / 'mage_2.png').convert('RGBA')
assert red_idle.size == static.size
assert red_idle.getchannel('A').tobytes() == static.getchannel('A').tobytes()
changed = 0
for p1, p2 in zip(static.getdata(), red_idle.getdata()):
    if p1 != p2:
        assert p1[3] == p2[3] and p1[3] > 0
        changed += 1
assert changed > 1000, 'Blue costume was not recolored'
print(f'Imported exact 736x736 P1 original and P2 recolored copy ({changed} RGB pixels changed)')
print('P1 GIF kept byte-for-byte; 4 full-resolution movement frames extracted')
